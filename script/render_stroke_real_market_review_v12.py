"""Render recorded evidence as figures, a workbook and a cited review document."""

from datetime import datetime, timezone
import json
from pathlib import Path
import re
from urllib.request import urlopen

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from matplotlib.ticker import MaxNLocator
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from review_stroke_real_markets_v12 import ROOT, bars, cold, file_record, save

OUT = ROOT / "output/bi_code_review/v12_real_market_audit"
FONT = FontProperties(fname="C:/Windows/Fonts/msyh.ttc")
plt.rcParams.update(
    {
        "font.family": FONT.get_name(),
        "axes.unicode_minus": False,
        "font.size": 10,
        "axes.spines.top": False,
        "axes.spines.right": False,
    }
)
BLUE, ORANGE, GRAY = "#245781", "#b85525", "#a6abb0"


def read(name):
    return json.loads((OUT / name).read_text(encoding="utf-8"))


def ranges(ax, candles, first, last):
    visible = [k for k in candles if first <= k.index <= last]
    for k in visible:
        ax.vlines(k.index, k.l, k.h, color=GRAY, lw=1.5, zorder=1)
        ax.hlines([k.l, k.h], k.index - 0.18, k.index + 0.18, color=GRAY, lw=1)
    ax.set_xlim(first - 0.8, last + 0.8)
    ax.xaxis.set_major_locator(MaxNLocator(nbins=9, integer=True))
    ax.grid(axis="y", alpha=0.16)
    ax.set_ylabel("记录价格")


def lines(ax, strokes, color, first, last, style="-", label=None):
    used = False
    for bi in strokes:
        a, b = bi["start"], bi["end"]
        if a["center"] < first or b["center"] > last:
            continue
        ax.plot(
            [a["center"], b["center"]],
            [a["value"], b["value"]],
            style,
            c=color,
            lw=2.1,
            zorder=3,
            label=label if not used else None,
        )
        used = True


def marked(ax, p, label, offset=13):
    ax.scatter(p["center"], p["value"], s=24, c="#20252b", zorder=5)
    ax.annotate(
        f"{label}：{p['center']} / {p['value']:g}",
        (p["center"], p["value"]),
        xytext=(0, offset),
        textcoords="offset points",
        ha="center",
        fontsize=9,
        va="bottom" if offset > 0 else "top",
    )


def savefig(fig, name):
    fig.savefig(OUT / (name + ".png"), dpi=180, facecolor="white", bbox_inches="tight")
    fig.savefig(OUT / (name + ".svg"), facecolor="white", bbox_inches="tight")
    plt.close(fig)


def retraction_chart(name, title, first, last):
    case = read("cases/" + name + ".json")
    frame = pd.read_parquet(case["input"]["path"])
    raw = bars(frame)
    middle, _ = cold(raw[: case["stage_raw_count"]])
    after, _ = cold(raw[: case["raw_count"]])
    fig, axs = plt.subplots(2, 1, figsize=(10.5, 6.1), sharex=True)
    points = (
        case["discovery"]["before_tail"][-3:] + case["discovery"]["after_tail"][-2:]
    )
    for ax, state, ks, color in (
        (axs[0], "middle", middle.cl_klines, BLUE),
        (axs[1], "after", after.cl_klines, ORANGE),
    ):
        ranges(ax, ks, first, last)
        lines(ax, case[state]["strokes"], color, first, last)
        ys = [k.l for k in after.cl_klines if first <= k.index <= last]
        zs = [k.h for k in after.cl_klines if first <= k.index <= last]
        span = max(zs) - min(ys)
        ax.set_ylim(min(ys) - span * 0.24, max(zs) + span * 0.24)
    for p, label in zip(points, "PABCD"):
        marked(axs[1], p, label, 12 if p["kind"] == "ding" else -12)
        if p["center"] <= points[-2]["center"]:
            marked(axs[0], p, label, 12 if p["kind"] == "ding" else -12)
    axs[0].set_title(title + "：C 出现后仍保留 P→A→B", loc="left", fontsize=12)
    axs[1].set_title(
        "D 出现后变为 P→C→D，原 P→A 的 is_done 曾为 True", loc="left", fontsize=11
    )
    axs[1].set_xlabel("包含处理后的 K 线中心序号；两幅图各自只使用当时已知行情")
    fig.tight_layout(h_pad=2.0)
    savefig(fig, name)


def boundary_chart():
    case = read("cases/recorded_SH.600189_5m_boundary.json")
    frame = pd.read_parquet(case["input"]["path"])
    merged, _ = cold(bars(frame.iloc[: case["raw_count"]]))
    fig, ax = plt.subplots(figsize=(10.5, 4.3))
    ranges(ax, merged.cl_klines, 1, 18)
    strokes = case["after"]["strokes"]
    lines(ax, strokes[:1], BLUE, 1, 18, label="保留的第一范围")
    lines(ax, strokes[1:], ORANGE, 1, 18, label="连接待定的后续范围")
    points = {p["center"]: p for p in case["points"]}
    for center, label in (
        (3, "P"),
        (8, "A"),
        (9, "B"),
        (10, "C"),
        (12, "D"),
        (17, "E"),
    ):
        p = points[center]
        marked(ax, p, label, 18 if p["kind"] == "ding" else -20)
    ax.set_ylim(7.255, 7.408)
    ax.set_xlabel("合并 K 线序号；A→D 虽相隔 4，区间内 C=7.37 高于 A=7.36")
    ax.set_title(
        "SH.600189 / 5m：两个旧起点均失去后续连接资格", loc="left", fontsize=12
    )
    ax.legend(loc="upper left", frameon=False, fontsize=9)
    fig.tight_layout()
    savefig(fig, "recorded_SH.600189_5m_boundary")


def initial_chart():
    item = read("initial_context_case.json")
    fig, axs = plt.subplots(3, 1, figsize=(10.5, 6.0), sharex=True, sharey=True)
    names = [
        "初始向上假设：顶分型成立",
        "初始向下假设：同一顶分型成立",
        "默认共同后缀：左肩被排除，顶分型漏掉",
    ]
    for ax, variant, title in zip(axs, item["variants"], names):
        for k in variant["first_candles"]:
            x = k["raw_center"]
            if x > 6:
                continue
            ax.vlines(
                x,
                k["low"],
                k["high"],
                color=BLUE if variant["assumed_initial_direction"] else ORANGE,
                lw=3,
            )
            ax.hlines([k["low"], k["high"]], x - 0.12, x + 0.12, color=GRAY)
            ax.annotate(
                "来源 " + ",".join(map(str, k["source_indices"])),
                (x, k["low"]),
                xytext=(0, -13),
                textcoords="offset points",
                ha="center",
                fontsize=9,
            )
        ax.set_title(title, loc="left", fontsize=11)
        ax.set_ylim(7.315, 7.42)
        ax.set_xlim(-0.5, 6.7)
        ax.grid(axis="y", alpha=0.15)
        if variant["assumed_initial_direction"]:
            ax.annotate(
                "共同顶：09:50 / 7.40",
                (3, 7.40),
                xytext=(4.1, 7.411),
                arrowprops={"arrowstyle": "->", "color": BLUE},
                fontsize=9,
            )
    axs[-1].set_xticks([0, 2, 3, 4, 6], ["09:35", "09:45", "09:50", "09:55", "10:05"])
    axs[-1].set_xlabel("2025-09-15，SH.600189 / 5m；假设分支用于核查，未修改生产配置")
    fig.tight_layout(h_pad=1.6)
    savefig(fig, "initial_context_consensus")


def tab(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(
            [
                json.dumps(v, ensure_ascii=False)
                if isinstance(v, (dict, list, tuple))
                else v
                for v in row
            ]
        )
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="28343D")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[letter].width = min(54, max(14, width + 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def workbook(sources):
    wb = Workbook()
    wb.remove(wb.active)
    rows = read("summary.json")
    ws = tab(
        wb,
        "样本汇总",
        [
            "数据集",
            "原始K数",
            "候选笔",
            "首个连续范围笔数",
            "未决连接",
            "线段",
            "首范围占候选比例",
            "行情末日",
            "来源路径",
            "SHA256",
        ],
        [
            [
                r["dataset"],
                r["quality"]["bars"],
                r["candidates"],
                r["first_contiguous"],
                r["boundaries"],
                r["segments"],
                f"=D{i}/C{i}",
                r["quality"]["end"],
                r["path"],
                r["sha256"],
            ]
            for i, r in enumerate(rows, 2)
        ],
    )
    for cell in list(ws.columns)[6][1:]:
        cell.number_format = "0.0%"
    verification = read("verification.json")
    tab(
        wb,
        "真实前缀撤回事件",
        [
            "数据集",
            "输入K数",
            "发生时刻",
            "撤回原is_done笔数",
            "原确认线段几何消失数",
            "逐行情前缀核实",
        ],
        [
            [
                e["dataset"],
                e["raw_count"],
                e["witness"],
                len(e["removed"]),
                len(e["removed_locked_segments"]),
                e["verified_primary_retraction"],
            ]
            for e in verification["primary_events"]
        ],
    )
    tab(
        wb,
        "处理模式",
        ["数据集", "增量批次数", "计算方式统计", "批量与增量相等"],
        [
            [r["dataset"], r["chunks"], r["mode_counts"], r["batch_incremental_equal"]]
            for r in verification["modes"]
        ],
    )
    tab(
        wb,
        "接续不可达证据",
        [
            "数据集",
            "可达起点中心",
            "起点价格",
            "首次破坏中心",
            "破坏时间",
            "破坏K高",
            "破坏K低",
            "局部合格后继FX序号",
        ],
        [
            [
                c["dataset"],
                v["point"]["center"],
                v["point"]["value"],
                v["first_price_invalidation"]["center"],
                v["first_price_invalidation"]["date"],
                v["first_price_invalidation"]["high"],
                v["first_price_invalidation"]["low"],
                v["valid_successors"],
            ]
            for c in read("boundary_certificates.json")
            for v in c["reachable"]
        ],
    )
    for n, p in enumerate(sorted((OUT / "cases").glob("*.json")), 1):
        case = json.loads(p.read_text(encoding="utf-8"))
        tab(
            wb,
            f"案例{n}_原始K",
            ["数据集", "序号", "时间", "开", "高", "低", "收", "量"],
            [
                [
                    case["dataset"],
                    r["index"],
                    r["date"],
                    r["open"],
                    r["high"],
                    r["low"],
                    r["close"],
                    r["volume"],
                ]
                for r in case["raw"]
            ],
        )
        tab(
            wb,
            f"案例{n}_合并K",
            ["数据集", "合并序号", "原始极值序号", "时间", "高", "低", "来源K序号"],
            [
                [
                    case["dataset"],
                    r["index"],
                    r["raw_center"],
                    r["date"],
                    r["high"],
                    r["low"],
                    r["raw_indices"],
                ]
                for r in case["merged"]
            ],
        )
        tab(
            wb,
            f"案例{n}_笔状态",
            [
                "时点",
                "起点中心",
                "终点中心",
                "起点时刻",
                "终点时刻",
                "起点价",
                "终点价",
                "范围",
                "is_done",
                "取舍时刻",
                "承接时刻",
                "待决",
            ],
            [
                [
                    state,
                    b["start"]["center"],
                    b["end"]["center"],
                    b["start"]["date"],
                    b["end"]["date"],
                    b["start"]["value"],
                    b["end"]["value"],
                    b["component"],
                    b["is_done"],
                    b["selected_at"],
                    b["continuation_at"],
                    b["selection_pending"],
                ]
                for state in ("middle", "before", "after")
                if state in case
                for b in case[state]["strokes"]
            ],
        )
    context = read("initial_context_case.json")
    tab(
        wb,
        "初始包含方向",
        [
            "假设方向",
            "合并序号",
            "原始极值序号",
            "时间",
            "高",
            "低",
            "来源原始序号",
            "最终范围笔数",
            "线段数",
        ],
        [
            [
                v["assumed_initial_direction"] or "默认共同后缀",
                k["center"],
                k["raw_center"],
                k["date"],
                k["high"],
                k["low"],
                k["source_indices"],
                v["components"],
                v["segments"],
            ]
            for v in context["variants"]
            for k in v["first_candles"]
        ],
    )
    scope = read("cases/recorded_SH.600189_5m_boundary.json")["windows"]
    tab(
        wb,
        "历史起点对照",
        [
            "起始原始序号",
            "起始时间",
            "K数",
            "候选笔",
            "各范围笔数",
            "首范围笔数",
            "线段数",
            "首范围结束",
        ],
        [
            [
                r[k]
                for k in (
                    "raw_start",
                    "start_time",
                    "bars",
                    "candidates",
                    "components",
                    "first_contiguous",
                    "segments",
                    "last_contiguous_time",
                )
            ]
            for r in scope
        ],
    )
    tab(wb, "Sources", ["编号", "来源类型", "文件路径", "定位", "SHA256"], sources)
    wb.save(OUT / "real_market_evidence.xlsx")


def hyperlink(paragraph, label, target):
    # Editor-style line anchors are retained in Markdown. Word opens the file.
    target = re.sub(r":\d+$", "", target)
    element = OxmlElement("w:hyperlink")
    relation = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    element.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "245781")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def rich(paragraph, text):
    parts = re.split(r"(\[[^\]]+\]\(<[^>]+>\)|\*\*[^*]+\*\*|`[^`]+`)", text)
    for item in parts:
        match = re.fullmatch(r"\[([^\]]+)\]\(<([^>]+)>\)", item)
        if match:
            hyperlink(paragraph, match[1], match[2])
        elif item.startswith("**"):
            paragraph.add_run(item[2:-2]).bold = True
        elif item.startswith("`"):
            run = paragraph.add_run(item[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(item)


def document(markdown):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.9)
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.18
    lines_ = markdown.splitlines()
    i = 0
    while i < len(lines_):
        line = lines_[i]
        if not line:
            i += 1
            continue
        if line.startswith("!["):
            path = re.search(r"\(<(.+)>\)", line)[1]
            doc.add_picture(path, width=Cm(16.9))
            i += 1
            continue
        if line.startswith("|"):
            data = []
            while i < len(lines_) and lines_[i].startswith("|"):
                cells = [x.strip() for x in lines_[i].strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", x) for x in cells):
                    data.append(cells)
                i += 1
            table = doc.add_table(rows=1, cols=len(data[0]))
            table.style = "Light Shading Accent 1"
            for j, c in enumerate(data[0]):
                rich(table.rows[0].cells[j].paragraphs[0], c)
            for row in data[1:]:
                cells = table.add_row().cells
                for j, c in enumerate(row):
                    rich(cells[j].paragraphs[0], c)
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        for run in p.runs:
                            run.font.size = Pt(8)
            continue
        p = doc.add_paragraph()
        rich(p, line)
        if i == 0:
            for run in p.runs:
                run.font.size = Pt(18)
        i += 1
    doc.save(OUT / "real_market_review.docx")


def main():
    retraction_chart("recorded_SH.600088_30m_retraction", "SH.600088 / 30m", 4, 23)
    retraction_chart("recorded_SH.600519_5m_retraction", "SH.600519 / 5m", 53, 75)
    boundary_chart()
    initial_chart()
    manifest = read("manifest.json")
    corpus = Path(manifest["sources"][0]["path"]).parent
    sources = []
    specifications = [
        (62, "正文", "37–118"),
        (65, "正文；注释另列", "31–100、166–214"),
        (66, "作者回复，非正文", "256–280"),
        (69, "正文", "28–49、151–184、205–223"),
        (77, "正文；注释另列", "118–148、169–295"),
        (64, "正文；注释另列", "73–103"),
        (70, "作者回复", "568–589、679–754"),
        (75, "作者回复", "565–580"),
    ]
    for n, (lesson, kind, locator) in enumerate(specifications, 1):
        path = next(corpus.glob(f"L{lesson:03d}_*.md"))
        # Actually reopen each cited original; no outside Chan text is used.
        path.read_text(encoding="utf-8")
        sources.append([n, kind, path.as_posix(), locator, file_record(path)["sha256"]])
    sources += [
        [
            "D1",
            "记录行情",
            str(ROOT / "output/playwright/center_coverage/deep_review/source_frames"),
            "六标的×三周期，详见样本汇总",
            None,
        ],
        [
            "D2",
            "测试行情",
            str(ROOT / "tests/fixtures"),
            "四份历史文件；SZ.002299_1m 另列稀疏性",
            None,
        ],
        [
            "E1",
            "复核证据",
            str(OUT / "verification.json"),
            "primary_events、modes",
            file_record(OUT / "verification.json")["sha256"],
        ],
    ]
    workbook(sources)
    docpath = ROOT / "docs/stroke_real_market_review_v12.md"
    markdown = docpath.read_text(encoding="utf-8")
    document(markdown)
    try:
        with urlopen("http://127.0.0.1:9900/readyz?market=a", timeout=5) as response:
            runtime = json.load(response)
    except Exception as exc:
        runtime = {"read_error": str(exc)}
    save(OUT / "runtime_readiness.json", runtime)
    for item in manifest["core_files"]:
        assert file_record(item["path"]) == item
    save(
        OUT / "delivery_manifest.json",
        {
            "created_at": datetime.now(timezone.utc),
            "core_unchanged": True,
            "source_files": [
                {"path": r[2], "locator": r[3], "sha256": r[4]} for r in sources[:8]
            ],
            "core_files": manifest["core_files"],
            "scripts": [
                file_record(ROOT / "script" / name)
                for name in (
                    "review_stroke_real_markets_v12.py",
                    "verify_stroke_real_market_cases_v12.py",
                    "certify_stroke_real_market_boundaries_v12.py",
                    "render_stroke_real_market_review_v12.py",
                )
            ],
            "report": file_record(docpath),
            "artifacts": [
                file_record(p)
                for p in sorted(OUT.glob("*"))
                if p.is_file()
                and p.name not in ("delivery_manifest.json", "manifest.json")
            ],
            "case_files": [
                file_record(p) for p in sorted((OUT / "cases").glob("*.json"))
            ],
            "production_modified": False,
            "runtime_restarted": False,
        },
    )
    print(
        "Created figures, real_market_evidence.xlsx, real_market_review.docx and delivery_manifest.json"
    )


if __name__ == "__main__":
    main()
