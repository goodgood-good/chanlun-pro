"""Render the recorded source-rule audit without invoking the stroke engine."""

import argparse
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]


def chart(data, folder):
    case = next(c for c in data["equal_reactivation"] if not c["mirror"] and c["endpoint_high"] == 16)
    prices = case["prices"]
    plt.rcParams.update({"font.family": "Microsoft YaHei", "font.size": 10, "axes.unicode_minus": False})
    fig, axes = plt.subplots(2, 1, figsize=(12.6, 8.7), sharex=True, sharey=True)
    fig.subplots_adjust(top=0.9, bottom=0.14, left=0.065, right=0.97, hspace=0.42)
    fig.suptitle("等价旧端点的排除条件改变后，是否重新审查", fontsize=17, x=0.065, ha="left")
    labels = {1: ("P", 20), 5: ("A", 6), 8: ("B", 14), 11: ("C", 6),
              15: ("D", 12), 17: ("F", 7), 19: ("E", 16)}
    navy, orange = "#2C5F83", "#B56A26"
    for ax, count, selected, title in zip(
        axes, (17, 21), ([1, 11, 15], [1, 11, 19]),
        ("前 17 根：A→D 的区间内有 B=14，高于 D=12；当前选择 P→C→D",
         "前 21 根：E=16 已越过 B=14；A→E 满足局部条件，当前仍选择 P→C→E"),
    ):
        ax.vlines(range(count), [p[1] for p in prices[:count]], [p[0] for p in prices[:count]],
                  color="#B5B5B5", linewidth=2, zorder=1)
        ax.plot(selected, [labels[i][1] for i in selected], color=navy, marker="o",
                linewidth=2.4, markersize=5.5, label="当前系统输出", zorder=3)
        if count == 21:
            ax.plot([1, 5, 19], [20, 6, 16], color=orange, linestyle="--", marker="o",
                    linewidth=2.2, markersize=5, label="需重新审查的较早等价候选 P→A→E", zorder=4)
        for i, (name, val) in labels.items():
            if i >= count - 1:
                continue
            below = name in ("A", "C", "F")
            offset = -20 if below else (27 if count == 21 and name == "D" else 10)
            ax.annotate(f"{name} ({i}, {val})", (i, val), xytext=(0, offset),
                        textcoords="offset points", ha="center", va="top" if below else "bottom",
                        fontsize=10, color="#242424", zorder=5)
        ax.set_title(title, loc="left", fontsize=11, pad=16)
        ax.set_ylim(3, 23)
        ax.set_xlim(-0.7, 20.7)
        ax.set_yticks([6, 10, 14, 18, 22])
        ax.set_ylabel("合成价格")
        ax.grid(axis="y", color="#E6E6E6", linewidth=0.7, zorder=0)
        ax.spines[["top", "right"]].set_visible(False)
        ax.spines[["left", "bottom"]].set_color("#C7C7C7")
        ax.tick_params(length=0, pad=6)
        ax.legend(loc="upper right", frameon=False, fontsize=9)
    axes[-1].set_xticks(range(21))
    axes[-1].set_xlabel("原始 K 线序号（从 0 开始；本例无包含、无缺口）", labelpad=12)
    fig.text(0.065, 0.05, "灰竖线为每根 K 线的高低区间。橙虚线是依据原文条件提出的待审候选，并非作者原图或已证明的唯一终局。",
             fontsize=9, color="#444444")
    fig.text(0.065, 0.023, "条件来源：第 62 课正文、第 66 课所附作者回复、第 77 课步骤二、三；原文路径及行号见配套工作簿 Sources。",
             fontsize=9, color="#444444")
    for suffix in ("png", "svg"):
        fig.savefig(folder / f"v10_equal_endpoint_reaudit.{suffix}", dpi=180, facecolor="white")
    plt.close(fig)


def workbook(data, folder, input_path):
    book = Workbook()
    book.remove(book.active)
    summary = book.create_sheet("Summary")
    summary.append(["案例", "镜像", "控制条件", "原始 K 数", "完整输入输出", "后段单独计算笔数",
                    "增量前缀数", "增量最终等于批量", "区间异常数", "再分异常数", "解释边界"])
    raw = book.create_sheet("Raw_Prices")
    raw.append(["案例", "序号（0 起）", "最高价", "最低价", "数据性质"])
    decisions = book.create_sheet("Selection")
    decisions.append(["案例", "分型中心", "性质", "价格", "局部合格前驱", "不可再分前驱", "处理前路径", "处理后路径",
                      "动作", "待接续", "本次候选父节点", "本次候选起源", "舍弃记录", "仍保留价格资格的起点"])
    transitions = book.create_sheet("Streaming")
    transitions.append(["案例", "已输入 K 数", "实际选中连接及完成状态", "待接续见证"])
    edges = book.create_sheet("Independent_Conditions")
    edges.append(["案例", "待核查路径", "起点中心", "终点中心", "局部条件成立", "同起止区间内最多合格连接数", "解释边界"])

    def compact(value):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))

    for group, cases in (("equal", data["equal_reactivation"]), ("origin", data["origin_wait"]),
                         ("initialization", data["initialization"]), ("completion", data["completion_retraction"])):
        for number, case in enumerate(cases):
            key = f"{group}_{number}"
            current = case.get("current", case.get("after"))
            strokes = current["strokes"]
            control = (f"E={case['endpoint_high']}" if group == "equal" else
                       f"extended={case['extended']}" if group == "origin" else
                       f"prior_context={case['context']}" if group == "initialization" else "11 根与 19 根比较")
            later = len(case["later_component_with_left_shoulder"]["strokes"]) if group == "origin" else None
            boundary = ("局部合格不等于已证明全图唯一；重新审查舍弃理由" if group == "equal" else
                        "后段局部合格，不等于可直接接入前缀" if group == "origin" else
                        "上下镜像是工程性质核验，原文未给无前情初始化默认值" if group == "initialization" else
                        "证明完成标记可撤换；不等于证明该标记就是原文完成")
            summary.append([key, case["mirror"], control, len(case["prices"]),
                            compact([[b["raw_start"], b["raw_end"], b["is_done"]] for b in strokes]),
                            later, case["streaming"]["prefixes"], case["streaming"]["final_matches_batch"],
                            len(current["range_audit"]), len(current["adjacency_audit"]), boundary])
            for index, (high, low) in enumerate(case["prices"]):
                raw.append([key, index, high, low, "为核查规则构造的合成数据；不是行情报价或原文图形"])
            for event in case.get("trace", []):
                decisions.append([key, event["center"], event["kind"], event["value"], compact(event["eligible"]),
                                  compact(event["irreducible"]), compact(event["previous"]), compact(event["selected"]),
                                  event["action"], event["pending"], event["candidate_parent"], event["candidate_origin"],
                                  compact(event["omitted"]), compact(event["live_starts"])])
            for event in case["streaming"]["changes"]:
                transitions.append([key, event["raw_bars"], compact(event["strokes"]), event["pending_continuation"]])
            for name in ("earlier_equal_alternative", "later_continuous_chain"):
                for edge in case.get(name, []):
                    edges.append([key, name, edge["start"], edge["end"], edge["locally_legal"],
                                  edge["longest_qualified_subdivision"], boundary])

    chart_data = book.create_sheet("Chart_Data")
    chart_data.append(["K 序号", "最高价", "最低价", "前 17 根面板可见", "分型标记", "分型价格", "前 17 根当前路径", "前 21 根当前路径", "待审候选路径"])
    case = next(c for c in data["equal_reactivation"] if not c["mirror"] and c["endpoint_high"] == 16)
    labels = {1: ("P", 20), 5: ("A", 6), 8: ("B", 14), 11: ("C", 6), 15: ("D", 12), 17: ("F", 7), 19: ("E", 16)}
    for index, (high, low) in enumerate(case["prices"]):
        label, price = labels.get(index, (None, None))
        chart_data.append([index, high, low, index < 17, label, price,
                           price if index in (1, 11, 15) else None,
                           price if index in (1, 11, 19) else None,
                           price if index in (1, 5, 19) else None])
    chart_data["B1"].comment = Comment("来源：本工作簿 Raw_Prices 的 equal_4；全部是合成高低价。", "Source")
    chart_data["I1"].comment = Comment("原文条件来源：Sources 的 L062、L066 作者回复、L077 步骤二、三；该路径为推导候选，不是原文逐字结论。", "Source")

    sources = book.create_sheet("Sources")
    sources.append(["来源类型", "文件/章节", "作者/来源", "定位", "路径（本地访问）", "SHA-256", "用途/边界"])
    for source in data["sources"]:
        filename = Path(source["path"]).name
        sources.append([source["kind"], filename, "缠中说禅；资料中的注释另行区分",
                        compact(source["line_spans"]), source["path"], source["sha256"],
                        "L066 是所附作者回复；L065/L077 的注释不作为作者正文。各案解释为本报告推导。"])
    for source in data["code"]:
        sources.append(["当前实现（不是原文依据）", Path(source["path"]).name, "工作区代码", "完整文件",
                        source["path"], source["sha256"], "定位实际执行的选择、状态和包含逻辑"])
    sources.append(["可复现实验数据", input_path.name, "规则核查合成案例", "16 组案例，846 个增量前缀",
                    str(input_path.resolve()), None, "图表与数字直接来自该 JSON；局部合法性不作为全局唯一性证明"])
    source_note = Comment("条件出处及实现证据见 Sources；所有价格都是合成实验输入。", "Source")
    for sheet in (summary, decisions, edges):
        sheet["A1"].comment = source_note

    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor="333333")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F2F2")
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(48, max(14, len(str(sheet.cell(1, column).value)) * 1.7))
    for col in ("M", "N"):
        decisions.column_dimensions[col].width = 65
    summary.column_dimensions["K"].width = 60
    transitions.column_dimensions["C"].width = 100
    sources.column_dimensions["E"].width = 90
    sources.column_dimensions["G"].width = 70
    path = folder / "v10_rule_reaudit_data.xlsx"
    book.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    assert check["Chart_Data"].max_row == 22 and check["Summary"].max_row == 17
    assert "Sources" in check.sheetnames
    check.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=ROOT / "output/bi_code_review/v10_rule_reaudit.json")
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8"))
    chart(data, args.input.parent)
    workbook(data, args.input.parent, args.input)
    print("Rendered PNG, SVG and supporting XLSX with Sources.")
